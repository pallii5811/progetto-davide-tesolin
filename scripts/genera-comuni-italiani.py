# -*- coding: utf-8 -*-
"""
Genera packages/core/src/shared/comuni-italiani-dati.ts dall'elenco ISTAT dei comuni.

Uso:
    python scripts/genera-comuni-italiani.py <Elenco-comuni-italiani.xlsx>

Il file si scarica da https://www.istat.it/classificazione/codici-dei-comuni-delle-province-e-delle-regioni/
(«Elenco dei codici e delle denominazioni delle unità territoriali»). Il foglio usato è quello
dei codici, il cui nome porta la data di aggiornamento: «CODICI al 21_02_2026».

Due fasi, come ogni generatore del progetto: prima si controlla TUTTO, poi si scrive. Se un
controllo fallisce il file TypeScript resta com'era, perché un elenco dei comuni sbagliato non
dà errori: fa semplicemente sparire una città dal selettore, o ne fa cercare un'altra.

Controlli:
- il foglio dei codici esiste e ha le colonne attese, cercate per NOME e non per posizione;
- più di 7.800 comuni (l'Italia ne ha circa 7.900: un numero molto più basso è un foglio sbagliato);
- codice catastale presente, nel formato lettera + tre cifre, e UNICO: è la chiave con cui il
  fornitore filtra (`townCode`), e un duplicato manderebbe una città a cercarne un'altra;
- sigla automobilistica di due lettere maiuscole, denominazione non vuota.
"""
import json
import pathlib
import re
import sys
import unicodedata

import openpyxl

RADICE = pathlib.Path(__file__).resolve().parent.parent
USCITA = RADICE / 'packages' / 'core' / 'src' / 'shared' / 'comuni-italiani-dati.ts'

COLONNA_NOME = 'Denominazione in italiano'
COLONNA_ALTRA_LINGUA = 'Denominazione altra lingua'
COLONNA_SIGLA = 'Sigla automobilistica'
COLONNA_CODICE = 'Codice Catastale del Comune'


def ferma(messaggio):
    print('FERMO:', messaggio)
    sys.exit(1)


def chiave_di_ordinamento(nome, sigla):
    senza_accenti = ''.join(
        c for c in unicodedata.normalize('NFD', nome) if unicodedata.category(c) != 'Mn'
    )
    return (senza_accenti.lower(), sigla)


def main():
    if len(sys.argv) != 2:
        ferma('indicare il percorso del file ISTAT (.xlsx)')
    sorgente = pathlib.Path(sys.argv[1])
    if not sorgente.is_file():
        ferma(f'file non trovato: {sorgente}')

    libro = openpyxl.load_workbook(sorgente, read_only=True, data_only=True)
    fogli = [n for n in libro.sheetnames if n.startswith('CODICI al ')]
    if len(fogli) != 1:
        ferma(f'atteso un solo foglio «CODICI al …», trovati: {libro.sheetnames}')
    foglio = fogli[0]
    righe = list(libro[foglio].iter_rows(values_only=True))
    intestazione = [str(v).strip() if v is not None else '' for v in righe[0]]

    indici = {}
    for colonna in (COLONNA_NOME, COLONNA_ALTRA_LINGUA, COLONNA_SIGLA, COLONNA_CODICE):
        if colonna not in intestazione:
            ferma(f'colonna mancante: «{colonna}»')
        indici[colonna] = intestazione.index(colonna)

    comuni = []
    for numero, riga in enumerate(righe[1:], start=2):
        if riga is None or all(v is None for v in riga):
            continue
        nome = str(riga[indici[COLONNA_NOME]] or '').strip()
        altra = str(riga[indici[COLONNA_ALTRA_LINGUA]] or '').strip()
        sigla = str(riga[indici[COLONNA_SIGLA]] or '').strip()
        codice = str(riga[indici[COLONNA_CODICE]] or '').strip()
        if nome == '':
            ferma(f'riga {numero}: denominazione vuota')
        if not re.fullmatch(r'[A-Z]{2}', sigla):
            ferma(f'riga {numero} ({nome}): sigla non valida «{sigla}»')
        if not re.fullmatch(r'[A-Z][0-9]{3}', codice):
            ferma(f'riga {numero} ({nome}): codice catastale non valido «{codice}»')
        comuni.append((nome, sigla, codice, altra if altra != '' and altra != nome else None))

    if len(comuni) < 7800:
        ferma(f'solo {len(comuni)} comuni: il foglio non è l\'elenco completo')

    visti = {}
    for nome, sigla, codice, _ in comuni:
        if codice in visti:
            ferma(f'codice catastale duplicato {codice}: {visti[codice]} e {nome} ({sigla})')
        visti[codice] = f'{nome} ({sigla})'

    comuni.sort(key=lambda c: chiave_di_ordinamento(c[0], c[1]))

    data = foglio.replace('CODICI al ', '').replace('_', '/')
    righe_ts = []
    for nome, sigla, codice, altra in comuni:
        valori = [nome, sigla, codice] + ([altra] if altra is not None else [])
        righe_ts.append('  [' + ', '.join(json.dumps(v, ensure_ascii=False) for v in valori) + '],')

    testo = (
        '/*\n'
        ' * GENERATO da scripts/genera-comuni-italiani.py: non si modifica a mano.\n'
        ' *\n'
        f' * Fonte: ISTAT, elenco dei codici dei comuni, foglio «{foglio}».\n'
        f' * {len(comuni)} comuni, ordinati per nome. Per ciascuno: denominazione in italiano, sigla\n'
        ' * automobilistica, codice catastale (Belfiore) e, dove esiste, la denominazione in altra lingua.\n'
        ' */\n'
        '\n'
        f'export const FONTE_COMUNI = {json.dumps(f"ISTAT, elenco dei comuni al {data}", ensure_ascii=False)};\n'
        '\n'
        'export const RIGHE_COMUNI: readonly (readonly [\n'
        '  nome: string,\n'
        '  sigla: string,\n'
        '  codiceCatastale: string,\n'
        '  altraLingua?: string,\n'
        '])[] = [\n' + '\n'.join(righe_ts) + '\n];\n'
    )

    USCITA.write_text(testo, encoding='utf-8', newline='\n')
    print(f'scritto {USCITA.relative_to(RADICE)}: {len(comuni)} comuni, '
          f'{sum(1 for c in comuni if c[3] is not None)} con denominazione in altra lingua, fonte «{foglio}»')


if __name__ == '__main__':
    main()
